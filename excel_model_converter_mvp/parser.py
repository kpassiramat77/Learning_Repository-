from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .models import (
    ExtractedFactor,
    ExtractedFormula,
    ExtractedLookupTable,
    ExtractedModel,
)


def parse_workbook(source: bytes | str | Path, filename: str | None = None) -> ExtractedModel:
    workbook = _load_workbook(source)
    workbook_name = filename or getattr(source, "name", None) or "workbook"

    factors: list[ExtractedFactor] = []
    lookup_tables: list[ExtractedLookupTable] = []
    formulas: list[ExtractedFormula] = []

    for sheet in workbook.worksheets:
        if _is_factor_sheet(sheet.title):
            factors.extend(_extract_factors(sheet))
        if _is_lookup_sheet(sheet.title):
            lookup_tables.extend(_extract_lookup_table(sheet))
        formulas.extend(_extract_formulas(sheet))

    return ExtractedModel(
        workbook_name=workbook_name,
        factors=factors,
        lookup_tables=lookup_tables,
        formulas=formulas,
    )


def _load_workbook(source: bytes | str | Path):
    if isinstance(source, (str, Path)):
        return load_workbook(source, data_only=False)
    return load_workbook(BytesIO(source), data_only=False)


def _is_factor_sheet(name: str) -> bool:
    lowered = name.lower()
    return "factor" in lowered or "input" in lowered


def _is_lookup_sheet(name: str) -> bool:
    lowered = name.lower()
    return "lookup" in lowered or "table" in lowered


def _extract_factors(sheet) -> list[ExtractedFactor]:
    rows = _non_empty_rows(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0][:2]]
    start_index = 1 if header == ["name", "value"] else 0
    factors: list[ExtractedFactor] = []
    for row in rows[start_index:]:
        if len(row) < 1 or row[0] is None:
            continue
        name = str(row[0]).strip()
        value = row[1] if len(row) > 1 else None
        factors.append(ExtractedFactor(name=name, value=value))
    return factors


def _extract_lookup_table(sheet) -> list[ExtractedLookupTable]:
    rows = _non_empty_rows(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    headers = [header or f"column_{index + 1}" for index, header in enumerate(headers)]
    data_rows = [list(row[: len(headers)]) for row in rows[1:]]
    return [
        ExtractedLookupTable(name=sheet.title, headers=headers, rows=data_rows)
    ]


def _extract_formulas(sheet) -> list[ExtractedFormula]:
    formulas: list[ExtractedFormula] = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f" or (
                isinstance(cell.value, str) and cell.value.startswith("=")
            ):
                formula = cell.value
                if isinstance(formula, str) and not formula.startswith("="):
                    formula = f"={formula}"
                formulas.append(
                    ExtractedFormula(
                        sheet=sheet.title, cell=cell.coordinate, formula=formula
                    )
                )
    return formulas


def _non_empty_rows(rows: Iterable[tuple]) -> list[tuple]:
    result: list[tuple] = []
    for row in rows:
        if any(cell is not None and str(cell).strip() != "" for cell in row):
            result.append(row)
    return result
