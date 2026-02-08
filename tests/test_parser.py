from openpyxl import Workbook

from excel_model_converter_mvp.parser import parse_workbook


def test_parse_workbook_extracts_sections(tmp_path):
    workbook_path = tmp_path / "model.xlsx"
    wb = Workbook()

    ws_inputs = wb.active
    ws_inputs.title = "Inputs"
    ws_inputs.append(["Name", "Value"])
    ws_inputs.append(["BaseRate", 0.1])

    ws_lookup = wb.create_sheet("Lookup_Table")
    ws_lookup.append(["Band", "Factor"])
    ws_lookup.append(["A", 1.0])
    ws_lookup.append(["B", 1.1])

    ws_model = wb.create_sheet("Model")
    ws_model["C2"] = "=A2*B2"

    wb.save(workbook_path)

    extracted = parse_workbook(workbook_path)

    assert extracted.factors
    assert extracted.factors[0].name == "BaseRate"
    assert extracted.lookup_tables
    assert extracted.lookup_tables[0].name == "Lookup_Table"
    assert extracted.lookup_tables[0].headers == ["Band", "Factor"]
    assert extracted.formulas
    assert extracted.formulas[0].formula.startswith("=")
