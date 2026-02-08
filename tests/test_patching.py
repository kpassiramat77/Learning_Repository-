from openpyxl import Workbook

from excel_model_converter_mvp.parser import parse_workbook
from excel_model_converter_mvp.patching import generate_patch
from excel_model_converter_mvp.service import create_default_orchestrator


def _build_workbook(path, formula: str) -> None:
    wb = Workbook()
    ws_inputs = wb.active
    ws_inputs.title = "Inputs"
    ws_inputs.append(["Name", "Value"])
    ws_inputs.append(["BaseRate", 0.1])

    ws_lookup = wb.create_sheet("Lookup_Table")
    ws_lookup.append(["Band", "Factor"])
    ws_lookup.append(["A", 1.0])

    ws_model = wb.create_sheet("Model")
    ws_model["C2"] = formula

    wb.save(path)


def test_generate_patch_for_single_formula_change(tmp_path):
    workbook_path_v1 = tmp_path / "model_v1.xlsx"
    workbook_path_v2 = tmp_path / "model_v2.xlsx"
    _build_workbook(workbook_path_v1, "=A2*B2")
    _build_workbook(workbook_path_v2, "=A2*B2+1")

    orchestrator = create_default_orchestrator()
    canonical_v1 = orchestrator.build_canonical(parse_workbook(workbook_path_v1))
    canonical_v2 = orchestrator.build_canonical(parse_workbook(workbook_path_v2))

    patch = generate_patch(canonical_v1, canonical_v2)

    assert patch.summary["replace"] == 1
    assert patch.summary["add"] == 0
    assert patch.summary["remove"] == 0
    assert patch.changes[0].previous_formula == "=A2*B2"
    assert patch.changes[0].value.raw_formula == "=A2*B2+1"
